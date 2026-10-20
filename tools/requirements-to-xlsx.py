#!/usr/bin/env python3
"""Export the requirements register from Markdown to an Excel workbook.

The Markdown is the single source of truth; this is only a presentation
transform, so there is no second copy of the data to keep in sync.

Usage:
    python tools/requirements-to-xlsx.py
Reads:  local/requirements.md
Writes: local/requirements.xlsx
"""

import os
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(REPO, "local", "requirements.md")
OUT = os.path.join(REPO, "local", "requirements.xlsx")

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF", size=10)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
L0_FILL = PatternFill("solid", fgColor="DDEBF7")
L1_FILL = PatternFill("solid", fgColor="FFF2CC")


def tables(text):
    """Yield each Markdown table in the file as a list of cell lists."""
    for block in re.split(r"\n\s*\n", text):
        rows = [ln for ln in block.splitlines() if ln.strip().startswith("|")]
        if len(rows) < 2:
            continue
        cells = [[c.strip() for c in ln.strip().strip("|").split("|")] for ln in rows]
        data = [r for r in cells if not all(re.fullmatch(r":?-{2,}:?", c or "-") for c in r)]
        if data:
            yield data


def main():
    text = open(SRC, encoding="utf-8").read()
    found = list(tables(text))
    if len(found) < 3:
        print("expected 3 tables in %s, found %d" % (SRC, len(found)))
        return 1
    assignment, derived, trace = found[0], found[1], found[2]

    # merge the two requirement tables and attach the traceability columns
    hdr = assignment[0] + ["Rationale", "Traced to"]
    tmap = {r[0]: r for r in trace[1:]}
    rows = []
    for src_rows, level in ((assignment, "L0"), (derived, "L1")):
        for r in src_rows[1:]:
            t = tmap.get(r[0], ["", "", ""])
            rows.append(r[:6] + [level] + [t[1], t[2]])

    headers = ["ID", "Requirement", "Src", "Verification", "Owner", "Status", "Level",
               "Rationale", "Traced to"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Requirements"
    ws.append(headers)
    for r in rows:
        ws.append(r)
    width = [10, 74, 16, 13, 22, 12, 6, 46, 22]
    for i, w in enumerate(width, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = HDR_FILL, HDR_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[1].height = 26
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
        level = ws.cell(row=row[0].row, column=7).value
        for cell in row:
            cell.border = BORDER
            cell.font = Font(size=10)
            cell.alignment = Alignment(wrap_text=cell.column in (2, 8), vertical="top")
            cell.fill = L0_FILL if level == "L0" else L1_FILL
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(headers)), ws.max_row)

    ws2 = wb.create_sheet("Traceability")
    for r in trace:
        ws2.append(r)
    for i, w in enumerate([10, 46, 22], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    for c in range(1, len(trace[0]) + 1):
        cell = ws2.cell(row=1, column=c)
        cell.fill, cell.font = HDR_FILL, HDR_FONT
    ws2.freeze_panes = "A2"

    ws3 = wb.create_sheet("Dashboard")
    ws3["A1"] = "Requirements dashboard"
    ws3["A1"].font = Font(bold=True, size=13, color="1F3864")
    n = ws3.max_row
    counts = [("Total", "=COUNTA(Requirements!A2:A%d)" % (len(rows) + 1)),
              ("From the assignment", '=COUNTIF(Requirements!G2:G%d,"L0")' % (len(rows) + 1)),
              ("Derived", '=COUNTIF(Requirements!G2:G%d,"L1")' % (len(rows) + 1)),
              ("Closed", '=COUNTIF(Requirements!F2:F%d,"Closed*")' % (len(rows) + 1)),
              ("In progress", '=COUNTIF(Requirements!F2:F%d,"In progress")' % (len(rows) + 1)),
              ("Open", '=COUNTIF(Requirements!F2:F%d,"Open")' % (len(rows) + 1))]
    for i, (label, formula) in enumerate(counts, start=3):
        ws3.cell(row=i, column=1, value=label).font = Font(size=11)
        ws3.cell(row=i, column=2, value=formula).font = Font(size=11)
    ws3.column_dimensions["A"].width = 24
    ws3.column_dimensions["B"].width = 10

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print("written: %s  (%d requirements)" % (OUT, len(rows)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
